"""Pruebas funcionales para contexto, aprobación, ediciones y regresión."""
import io
import zipfile
import pytest
from fastapi.testclient import TestClient
from docx import Document

KEY='sky-v96-owner-key-secure-and-local-only'
HEAD={'X-SKY-KEY':KEY}

@pytest.fixture
def environment(tmp_path,monkeypatch):
    monkeypatch.setenv('SKY_API_KEY',KEY)
    monkeypatch.setenv('SKY_WORKFLOWS_ENABLED','0')
    for key,name in [('SKY_INTELLIGENT_DB','intelligent.db'),('SKY_WORKFLOWS_DB','workflows.db'),
                     ('SKY_PROJECTS_DB','projects.db'),('SKY_WORKBENCH_DB','workbench.db'),
                     ('SKY_METAVERSE_DB','worlds.db'),('SKY_COMMUNITY_DB','community.db'),
                     ('SKY_TELEMETRY_DB','telemetry.db')]:
        monkeypatch.setenv(key,str(tmp_path/name))
    from neural_core import sky_workbench
    monkeypatch.setattr(sky_workbench,'STORE',tmp_path)
    from metaverse.api import app
    with TestClient(app) as client:
        yield client,tmp_path


def proposal(c,message):
    r=c.post('/api/sky/intelligent/message',headers=HEAD,json={'message':message})
    assert r.status_code==200,r.text
    return r.json()['request']


def execute(c,stages_message):
    p=proposal(c,stages_message)
    assert c.post(f'/api/sky/intelligent/requests/{p["id"]}/approve',headers=HEAD,
                  json={'approve':True,'expected_fingerprint':p['fingerprint']}).status_code==200
    from neural_core.sky_workflows import run_once
    for _ in p['stages']:run_once()
    d=c.get(f'/api/sky/intelligent/requests/{p["id"]}',headers=HEAD).json()
    assert d['status']=='completed',d
    return d


def test_intent_variations_and_avoid_questions():
    from neural_core.sky_intelligent import creation_intent
    for m in ['Elaboráme una presentación de historia', 'Quiero que crees un proyecto educativo completo',
              'Necesito crear un documento Word de turismo', 'Redactame un informe cultural']:
        assert creation_intent(m),m
    for m in ['¿Podés hacer una presentación?', 'Cómo puedo crear un mundo', 'Cuál es el mejor curso']:
        assert not creation_intent(m),m


def test_plan_revision_removes_world_and_old_fingerprint_invalid(environment):
    c,_=environment
    initial=proposal(c,'Crea mi proyecto completo de academia educativa y presentación')
    assert len(initial['stages'])==4
    url=f'/api/sky/intelligent/requests/{initial["id"]}'
    assert c.post(url+'/revise-plan',json={'instruction':'Quita el mundo 3D del proyecto','expected_fingerprint':initial['fingerprint']}).status_code==401
    r=c.post(url+'/revise-plan',headers=HEAD,json={'instruction':'Quita el mundo 3D del proyecto','expected_fingerprint':initial['fingerprint']})
    assert r.status_code==200,r.text
    edited=r.json()
    assert len(edited['stages'])==3, edited['stages']
    assert edited['fingerprint']!=initial['fingerprint']
    assert c.post(url+'/approve',headers=HEAD,json={'approve':True,'expected_fingerprint':initial['fingerprint']}).status_code==409
    assert c.get('/api/sky/projects',headers=HEAD).json()==[]
    assert c.post(url+'/revise-plan',headers=HEAD,json={'instruction':'Agrega ejemplos para docentes','expected_fingerprint':initial['fingerprint']}).status_code==409


def test_contextual_document_genuine_and_traceability(environment):
    c,_=environment
    result=execute(c,'Elabora una presentación PowerPoint de cultura para docentes y estudiantes')
    assert [s['action'] for s in result['stages']]==['plan','document']
    entry=next(x for x in result['deliverables'] if x['kind']=='file')
    binary=c.get(entry['download'],headers=HEAD).content
    with zipfile.ZipFile(io.BytesIO(binary)) as presentation:
        assert any(x.startswith('ppt/slides/slide') for x in presentation.namelist())
        combined=' '.join(presentation.read(x).decode('utf-8') for x in presentation.namelist() if x.startswith('ppt/slides/slide') and x.endswith('.xml'))
        assert 'docentes' in combined
    assert result['dependencies'][1]['depends_on']==[]


def test_tabular_excel_requires_actual_data(environment):
    c,_=environment
    no_data=c.post('/api/sky/intelligent/message',headers=HEAD,json={'message':'Genera un Excel de ingresos del año'})
    assert no_data.status_code==422
    content='Genera un Excel con datos proporcionados: ```csv\nMes,Ingresos,Gastos\nEnero,100,70\nFebrero,140,80\n```'
    d=execute(c,content)
    file=next(x for x in d['deliverables'] if x['kind']=='file')
    binary=c.get(file['download'],headers=HEAD).content
    assert binary[:2]==b'PK'
    from openpyxl import load_workbook
    w=load_workbook(io.BytesIO(binary),read_only=True)
    texts=[str(cell) for row in w.active.iter_rows(values_only=True) for cell in row if cell is not None]
    assert 'Enero' in texts and 'Febrero' in texts


def test_artifact_revision_approval_and_preserve_original(environment):
    c,_=environment
    d=execute(c,'Crea un documento Word sobre historia cultural para estudiantes')
    original=next(x for x in d['deliverables'] if x['kind']=='file')
    old=c.get(original['download'],headers=HEAD).content
    endpoint=f'/api/sky/intelligent/requests/{d["id"]}/revisions'
    assert c.post(endpoint,json={'source_item_id':original['item_id'],'instruction':'Agregar una sección de actividades educativas'}).status_code==401
    r=c.post(endpoint,headers=HEAD,json={'source_item_id':original['item_id'],'instruction':'Agregar una sección de actividades educativas'})
    assert r.status_code==200,r.text
    revision=r.json()
    assert revision['status']=='awaiting_approval'
    assert c.post(f'/api/sky/intelligent/revisions/{revision["id"]}/approve',headers=HEAD,
                  json={'approve':True,'expected_fingerprint':'0'*64}).status_code==409
    assert c.get(original['download'],headers=HEAD).content==old
    confirmed=c.post(f'/api/sky/intelligent/revisions/{revision["id"]}/approve',headers=HEAD,
                  json={'approve':True,'expected_fingerprint':revision['fingerprint']})
    assert confirmed.status_code==200,confirmed.text
    assert confirmed.json()['status']=='completed'
    refreshed=c.get(f'/api/sky/intelligent/requests/{d["id"]}',headers=HEAD).json()
    assert len(refreshed['revisions'])==1
    assert refreshed['revisions'][0]['status']=='completed'
    revised_output=refreshed['revisions'][0]['result']
    changed=c.get(revised_output['download'],headers=HEAD).content
    assert changed!=old
    text=' '.join(p.text for p in Document(io.BytesIO(changed)).paragraphs)
    assert 'actividades educativas' in text
    assert c.get(original['download'],headers=HEAD).content==old
    assert c.post(f'/api/sky/intelligent/revisions/{revision["id"]}/approve',headers=HEAD,
                  json={'approve':True,'expected_fingerprint':revision['fingerprint']}).status_code==409
    manifest=c.get(refreshed['project_export'],headers=HEAD).content
    with zipfile.ZipFile(io.BytesIO(manifest)) as z:
        assert sum(n.endswith('.docx') for n in z.namelist())==2


def test_revision_cannot_access_unlinked_item_or_community(environment):
    c,_=environment
    first=execute(c,'Crea documento Word de astronomía para docentes')
    second=execute(c,'Crea documento Word de gastronomía para estudiantes')
    foreign=next(x for x in second['deliverables'] if x['kind']=='file')
    end=f'/api/sky/intelligent/requests/{first["id"]}/revisions'
    assert c.post(end,headers=HEAD,json={'source_item_id':foreign['item_id'],'instruction':'Agregar tablas y conclusiones'}).status_code==403
    assert c.post(end,headers={'Authorization':'Bearer user'},json={'source_item_id':foreign['item_id'],'instruction':'Agregar tablas y conclusiones'}).status_code==401


def test_revision_world_and_design_are_new_versions(environment):
    c,_=environment
    d=execute(c,'Crea un proyecto completo de ciudad virtual y diseño web')
    for kind in ('world','experience'):
        source=next(x for x in d['deliverables'] if x['kind']==kind)
        resp=c.post(f'/api/sky/intelligent/requests/{d["id"]}/revisions',headers=HEAD,
                    json={'source_item_id':source['item_id'],'instruction':'Agregar un ambiente de naturaleza nocturna'})
        assert resp.status_code==200,resp.text
        revision=resp.json()
        final=c.post(f'/api/sky/intelligent/revisions/{revision["id"]}/approve',headers=HEAD,
                     json={'approve':True,'expected_fingerprint':revision['fingerprint']})
        assert final.status_code==200,final.text
        from neural_core.sky_workbench import get_item
        new=get_item(final.json()['result_item_id'])
        assert new['status']=='completada'
        assert new['id']!=source['item_id']


def test_minimal_ui_retained_and_creator_only(environment):
    c,_=environment
    html=c.get('/sky').text
    assert '<section class="welcome">' in html
    assert 'BAQUEANO' not in html
    script=c.get('/static/sky-home.js').text
    assert 'Editar este plan' in script and 'Preparar revisión' in script
    assert c.get('/api/sky/intelligent/requests',headers={'Authorization':'Bearer community'}).status_code==401
    assert c.get('/api/sky/intelligent/requests',headers=HEAD).status_code==200
