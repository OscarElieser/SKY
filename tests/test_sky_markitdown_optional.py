"""Regression and optional live conversion checks for SKY MarkItDown preview."""
from __future__ import annotations

import io
import json
import os
import subprocess
from types import SimpleNamespace

import pytest

from neural_core import sky_markitdown_optional as markitdown


def test_feature_off_by_default(monkeypatch):
    monkeypatch.delenv("SKY_MARKITDOWN_ENABLED", raising=False)
    info = markitdown.readiness()
    assert info["enabled"] is False
    assert info["automatically_indexed"] is False
    with pytest.raises(markitdown.MarkItDownUnavailable):
        markitdown.preview("test.pdf", b"%PDF-1.4", approved=True)


def test_explicit_approval_and_allowlist(monkeypatch):
    monkeypatch.setenv("SKY_MARKITDOWN_ENABLED", "1")
    with pytest.raises(PermissionError):
        markitdown.preview("test.docx", b"PKhello")
    for name in ("../../secret.docx", "..\\secret.docx", "evil.zip", "evil.mp4"):
        with pytest.raises(ValueError):
            markitdown.preview(name, b"PKhello", approved=True)
    with pytest.raises(ValueError):
        markitdown.preview("readme.pdf", b"PKhello", approved=True)
    with pytest.raises(ValueError):
        markitdown.preview("example.docx", b"not-a-zip", approved=True)
    with pytest.raises(ValueError):
        markitdown.preview("empty.pdf", b"", approved=True)
    with pytest.raises(ValueError):
        markitdown.preview("huge.pdf", b"%PDF-" + b"a" * markitdown.MAX_INPUT_BYTES, approved=True)


def test_child_process_receives_no_secrets(monkeypatch):
    monkeypatch.setenv("SKY_MARKITDOWN_ENABLED", "1")
    monkeypatch.setenv("OPENAI_API_KEY", "MUST_NOT_LEAK")
    monkeypatch.setenv("SKY_NEURAL_MEMORY_KEY", "MUST_NOT_LEAK")
    called = []
    def fake_run(args, **kwargs):
        called.append((args, kwargs))
        return SimpleNamespace(
            returncode=0,
            stdout=json.dumps({"markdown": "# Hola SKY", "truncated": False}).encode(),
        )
    monkeypatch.setattr(subprocess, "run", fake_run)
    r = markitdown.preview("manual.docx", b"PKplaceholder", approved=True)
    assert r["markdown"] == "# Hola SKY"
    assert r["indexed"] is False and r["model_weights_updated"] is False
    assert r["source_trust"] == "owner_uploaded_unverified"
    args, kwargs = called[0]
    assert args[1:3] == ["-m", "neural_core.sky_markitdown_worker"]
    assert args[3] == ".docx"
    assert kwargs["input"] == b"PKplaceholder"
    assert kwargs["timeout"] <= 30
    assert kwargs["stderr"] == subprocess.DEVNULL
    assert not any("KEY" in key or "TOKEN" in key for key in kwargs["env"])


def test_converter_failures_report_unavailable(monkeypatch):
    monkeypatch.setenv("SKY_MARKITDOWN_ENABLED", "1")
    monkeypatch.setattr(subprocess, "run", lambda *a, **k: SimpleNamespace(returncode=3, stdout=b""))
    with pytest.raises(markitdown.MarkItDownUnavailable):
        markitdown.preview("manual.pdf", b"%PDF-1.4", approved=True)
    def too_slow(*args, **kwargs):
        raise subprocess.TimeoutExpired(cmd=args[0], timeout=kwargs["timeout"])
    monkeypatch.setattr(subprocess, "run", too_slow)
    with pytest.raises(markitdown.MarkItDownUnavailable, match="tiempo"):
        markitdown.preview("manual.pdf", b"%PDF-1.4", approved=True)


def test_preview_api_is_owner_protected():
    # Import the app first: importing knowledge_api directly would create a
    # circular import with metaverse.api's router registration.
    from metaverse.api import app
    matches = [route for route in app.routes
               if getattr(route, "path", None) == "/api/sky/knowledge/convert-markdown-preview"]
    assert len(matches) == 1
    assert any(dep.call.__name__ == "require_key"
               for dep in matches[0].dependant.dependencies)


def test_real_docx_conversion_when_optional_package_installed(monkeypatch):
    pytest.importorskip("markitdown")
    from docx import Document
    doc = Document()
    doc.add_heading("Prueba SKY Knowledge", level=1)
    doc.add_paragraph("Contenido de prueba para Markdown.")
    buffer = io.BytesIO()
    doc.save(buffer)
    monkeypatch.setenv("SKY_MARKITDOWN_ENABLED", "1")
    result = markitdown.preview("clase.docx", buffer.getvalue(), approved=True)
    assert "Prueba SKY Knowledge" in result["markdown"]
    assert "Contenido de prueba" in result["markdown"]
    assert result["indexed"] is False
    assert result["content_extracted"] is True
