"""Se prueban los archivos REALMENTE exportados y la API autenticada."""
from io import BytesIO
import pytest
from fastapi.testclient import TestClient
from neural_core.sky_office import create_document
from pptx import Presentation
from docx import Document
from pypdf import PdfReader
from artifact_tool import Blob,SpreadsheetFile

KEY='skys-personal-studio-authentication-key-321'

@pytest.fixture()
def client(monkeypatch,tmp_path):
    monkeypatch.setenv('SKY_API_KEY',KEY)
    monkeypatch.setenv('SKY_METAVERSE_DB',str(tmp_path/'worlds.sqlite3'))
    monkeypatch.setenv('SKY_COMMUNITY_DB',str(tmp_path/'community.sqlite3'))
    from metaverse.api import app
    with TestClient(app) as c:yield c


def sample(kind):return create_document({'kind':kind,'title':'Capacitación en innovación','content':'# Objetivos\n- Diseñar proyectos sostenibles\n# Actividades\nInvestigar turismo en Nicaragua'})

def test_word_is_editable_document():
    result=sample('docx');doc=Document(BytesIO(result.data))
    assert 'Capacitación en innovación' in '\n'.join(p.text for p in doc.paragraphs)
    assert any('Objetivos' in p.text for p in doc.paragraphs)


def test_powerpoint_has_real_editable_slides():
    deck=Presentation(BytesIO(sample('pptx').data))
    assert len(deck.slides)>=3
    assert any('Diseñar proyectos sostenibles' in shape.text for slide in deck.slides for shape in slide.shapes if shape.has_text_frame)


def test_pdf_is_actual_extractable_pdf():
    doc=PdfReader(BytesIO(sample('pdf').data))
    assert len(doc.pages)>=1
    assert 'innovación' in doc.pages[0].extract_text()


def test_xlsx_data_and_safe_formulas(tmp_path):
    result=create_document({'kind':'xlsx','title':'Reporte de ingresos','headers':['Concepto','Monto'],
         'rows':[['Cursos',350],['=SUM(A1:A9)',0]]})
    path=tmp_path/'r.xlsx';path.write_bytes(result.data)
    book=SpreadsheetFile.import_xlsx(Blob.load(str(path)))
    tab=book.worksheets.get_item_at(0)
    assert tab.get_range('A3:B5').values[0]==['Concepto','Monto']
    assert tab.get_range('B4').values==[[350]]
    assert tab.get_range('A5').values[0][0]=='=SUM(A1:A9)'
    from zipfile import ZipFile
    with ZipFile(BytesIO(result.data)) as z:
        xml=z.read('xl/worksheets/sheet1.xml').decode()
    assert '<x:c r="A5" t="str"' in xml and '<x:f>' not in xml


def test_budget_template_contains_actual_formulas():
    result=create_document({'kind':'xlsx','title':'Presupuesto','template':'presupuesto'})
    assert result.data.startswith(b'PK')
    from zipfile import ZipFile
    with ZipFile(BytesIO(result.data)) as z:
        worksheet=z.read('xl/worksheets/sheet1.xml').decode()
    assert 'SUM(' in worksheet and 'B4-C4' in worksheet


def test_csv_formula_injection_blocked():
    result=create_document({'kind':'csv','title':'Test','headers':['Dato'], 'rows':[['=1+2'],['@cmd']]})
    text=result.data.decode('utf-8-sig')
    assert "'=1+2" in text and "'@cmd" in text


def test_api_auth_and_real_file_generation(client):
    req={'kind':'pptx','title':'SKY Educación','content':'# Tema\nEl valor de la innovación'}
    assert client.post('/api/office/generate',json=req).status_code==401
    response=client.post('/api/office/generate',json=req,headers={'X-SKY-KEY':KEY})
    assert response.status_code==200,response.text[:300] if response.status_code!=200 else ''
    assert response.content[:2]==b'PK' and 'presentationml' in response.headers['content-type']
    assert len(Presentation(BytesIO(response.content)).slides)==2


def test_api_rejects_unknown_format(client):
    r=client.post('/api/office/generate',json={'kind':'exe','title':'No'},headers={'X-SKY-KEY':KEY})
    assert r.status_code==422


def test_api_office_page_visible_and_formats_private(client):
    assert client.get('/office').status_code==200
    assert client.get('/api/office/formats').status_code==401
    assert client.get('/api/office/formats',headers={'X-SKY-KEY':KEY}).status_code==200
