"""SKY multimodal reader: real extraction, privacy and safe limits."""
import io
import os
from pathlib import Path

import pytest
from fastapi.testclient import TestClient


@pytest.fixture
def client(tmp_path, monkeypatch):
    monkeypatch.setenv("SKY_API_KEY","sky-multimodal-owner-secret-key-for-tests-2026")
    monkeypatch.setenv("SKY_WORKFLOWS_ENABLED","0")
    monkeypatch.setenv("SKY_CREATIVE_DIR",str(tmp_path/"creative"))
    from metaverse import api
    monkeypatch.setattr(api,"UPLOADS_DIR",tmp_path/"private_media")
    with TestClient(api.app) as value: yield value


def owner():
    return {"X-SKY-KEY":"sky-multimodal-owner-secret-key-for-tests-2026"}


def test_docx_excel_pptx_pdf_real_local_extraction():
    from neural_core.sky_multimodal_reader import analyze
    from docx import Document
    doc=Document()
    doc.add_heading("SKY Docente",level=1)
    doc.add_paragraph("Plan de clases verificable")
    b=io.BytesIO();doc.save(b)
    a=analyze("clases.docx",b.getvalue())
    assert a["content_extracted"] and "Plan de clases" in a["text"]
    assert a["learning_performed"] is False

    from openpyxl import Workbook
    wb=Workbook();wb.active.title="Ventas";wb.active.append(["Producto","Precio"]);wb.active.append(["Mapa",27])
    buf=io.BytesIO();wb.save(buf)
    a=analyze("plan.xlsx",buf.getvalue())
    assert "Ventas" in a["text"] and "Mapa" in a["text"]
    assert "fórmulas no se ejecutan" in a["limitations"][0]

    from pptx import Presentation
    p=Presentation();slide=p.slides.add_slide(p.slide_layouts[5])
    slide.shapes.title.text="Presentación SKY"
    buf=io.BytesIO();p.save(buf)
    a=analyze("presentacion.pptx",buf.getvalue())
    assert a["metadata"]["slides"]==1
    assert "Presentación SKY" in a["text"]

    from reportlab.pdfgen.canvas import Canvas
    buf=io.BytesIO();c=Canvas(buf);c.drawString(50,700,"Informe SKY comprobado");c.showPage();c.save()
    a=analyze("informe.pdf",buf.getvalue())
    assert "Informe SKY comprobado" in a["text"]


def test_image_metadata_not_false_visual_understanding():
    from neural_core.sky_multimodal_reader import analyze
    from PIL import Image
    b=io.BytesIO();Image.new("RGB",(64,32)).save(b,format="PNG")
    result=analyze("imagen.png",b.getvalue())
    assert result["metadata"]["width"]==64 and result["metadata"]["height"]==32
    assert result["visual_semantics_verified"] is False
    assert result["content_extracted"] is False
    assert any("modelo multimodal" in x for x in result["limitations"])


def test_file_safety_rejects_unknown_and_oversized():
    from neural_core.sky_multimodal_reader import analyze
    with pytest.raises(ValueError): analyze("shell.exe",b"MZ")
    with pytest.raises(ValueError): analyze("../private.py",b"x"*(15*1024*1024+1))
    with pytest.raises(ValueError): analyze("fake.pdf",b"this-is-not-pdf")
    with pytest.raises(ValueError): analyze("bad.txt",b"")


def test_authenticated_reader_does_not_remember(client, tmp_path):
    data=b"# Programa tecnico\nProyecto de SKY"
    unauthed=client.post("/api/sky/reader/analyze",files={"file":("plan.md",data,"text/markdown")})
    assert unauthed.status_code==401
    res=client.post("/api/sky/reader/analyze",headers=owner(),
                    files={"file":("plan.md",data,"text/markdown")})
    assert res.status_code==200,res.text
    obj=res.json()
    assert obj["content_extracted"] is True
    assert "Proyecto de SKY" in obj["text"]
    assert obj["learning_performed"] is False
    assert not (tmp_path/"private_media").exists()
    assert client.get("/api/sky/reader/capabilities",headers=owner()).json()["full_video_understanding"] is False


def test_old_upload_no_longer_claims_assimilation(client,tmp_path):
    from neural_core import storage
    data=b"Las fuentes no se asimilan automaticamente"
    resp=client.post("/api/chat/upload",headers=owner(),
                     files={"file":("example.txt",data,"text/plain")})
    assert resp.status_code==200,resp.text
    assert resp.json()["learning_performed"] is False
    assert resp.json()["doc_id"] is None
    assert data.decode() in resp.json()["full_text"]


def test_private_media_requires_short_lived_url_signature(client):
    from PIL import Image
    f=io.BytesIO();Image.new("RGB",(25,19)).save(f,format="PNG")
    response=client.post("/api/chat/upload",headers=owner(),
                         files={"file":("sample.png",f.getvalue(),"image/png")})
    assert response.status_code==200,response.text
    url=response.json()["media_url"]
    assert url and "sig=" in url
    filename=url.split("?")[0].split("/")[-1]
    assert client.get(f"/api/media/uploads/{filename}").status_code==403
    real=client.get(url)
    assert real.status_code==200 and real.content.startswith(b"\x89PNG")
    assert client.get(url+"&sig=bad").status_code in (403,422)
