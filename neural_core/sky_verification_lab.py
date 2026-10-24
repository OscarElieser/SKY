"""SKY Verification Lab — verificaciones observables, evidencia y límites explícitos.

Complementa a SKY CORE sin levantar otra API, duplicar el agente o conceder permisos.
No transmite documentos a terceros, no persiste respuestas de estudiantes y no
asegura que una operación fue completada sólo porque se envió.
"""
from __future__ import annotations

import hashlib
import importlib.util
import io
import os
import platform
import re
import socket
from pathlib import Path
from urllib.parse import urlsplit
from urllib.request import Request, build_opener, HTTPRedirectHandler, HTTPSHandler


MAX_DOC_BYTES = 5_000_000
MAX_PAGES = 60
MAX_EXCERPT = 2_000
MAX_TOTAL_CHARS = 32_000
MODES = {"pdf", "docx", "txt", "md", "rst"}


def capability_catalogue() -> dict:
    """Contraste honesto del video: arquitectura propia, no paridad con terceros."""
    from . import sky_subagents, sky_pro_skills
    return {
        "central_agent": "SKY CORE",
        "single_existing_api": True,
        "subagents": sky_subagents.catalogue(),
        "professional_guides": len(sky_pro_skills.catalog()),
        "skills_are_installed_external_products": False,
        "mcp_general_client_implemented": False,
        "mcp_note": "El catálogo de conectores no implementa el protocolo MCP ni valida sesiones externas.",
        "security": "SKY SENTINEL y políticas existentes: no se modifican por este módulo.",
        "external_services_need_explicit_credentials": True,
        "real_windows_hardware_tested_here": False,
        "features": ["runtime_readiness", "authorized_connectors_catalogue",
                     "page_based_evidence", "claim_crosscheck_candidates",
                     "formative_review_without_grading", "existing_kpi_summary",
                     "non_executing_subagent_plan"],
    }


def _found_module(name: str) -> bool:
    try:
        return importlib.util.find_spec(name) is not None
    except (ImportError, ValueError):
        return False


def _loopback_url(url: str) -> bool:
    try:
        p = urlsplit(url)
        return (p.scheme == "http" and p.hostname in {"127.0.0.1", "localhost", "::1"}
                and p.username is None and p.password is None
                and p.port == 11434 and p.path in {"", "/", "/api/tags"}
                and not p.query and not p.fragment)
    except ValueError:
        return False


class _NoRedirect(HTTPRedirectHandler):
    def redirect_request(self, req, fp, code, msg, headers, newurl):
        return None


def runtime_readiness(probe_ollama: bool = False) -> dict:
    """Inspección local no intrusiva. Un puerto abierto NO acredita autenticación."""
    stt = {name: _found_module(name) for name in ("vosk", "faster_whisper", "sounddevice")}
    piper_bin = os.getenv("SKY_PIPER_BIN", "").strip()
    piper_voice = os.getenv("SKY_PIPER_VOICE", "").strip()
    piper = {
        "binary_declared_and_exists": bool(piper_bin and Path(piper_bin).is_file()),
        "voice_declared_and_exists": bool(piper_voice and Path(piper_voice).is_file()),
        "audio_playback_verified": False,
    }
    ollama_url = os.getenv("SKY_OLLAMA_URL", "http://127.0.0.1:11434").rstrip("/")
    probe = {"attempted": False, "status": "not_checked", "model_available": None}
    if probe_ollama:
        if not _loopback_url(ollama_url):
            probe = {"attempted": False, "status": "blocked_non_loopback_or_nonstandard_port",
                     "model_available": None}
        else:
            try:
                # Sólo una petición local de bajo riesgo; sin redirecciones, tokens o acciones.
                opener = build_opener(_NoRedirect(), HTTPSHandler())
                with opener.open(Request(ollama_url + "/api/tags", headers={"Accept": "application/json"}),
                                 timeout=1.5) as response:
                    payload = response.read(65_537)
                if len(payload) > 65_536:
                    probe = {"attempted": True, "status": "response_too_large", "model_available": None}
                else:
                    import json
                    parsed = json.loads(payload)
                    models = parsed.get("models", []) if isinstance(parsed, dict) else []
                    probe = {"attempted": True, "status": "http_ok" if isinstance(models, list) else "invalid_response",
                             "model_available": bool(models) if isinstance(models, list) else None}
            except (OSError, ValueError, TypeError, TimeoutError):
                probe = {"attempted": True, "status": "unreachable_or_invalid", "model_available": None}
    device_policy = Path.home() / ".neuralt" / "device_agent_policy.json"
    return {
        "platform": platform.system(),
        "stt_dependencies_detected": stt,
        "microphone_recording_tested": False,
        "piper": piper,
        "ollama": probe,
        "device_policy_file_exists": device_policy.is_file(),
        "agent_authenticated_action_tested": False,
        "permissions_hud_agent_unified_verified": False,
        "note": "Estas son señales de preparación, no pruebas físicas de voz, hardware ni acciones Windows.",
    }


def connector_catalogue() -> dict:
    """Sólo detecta presencia de configuración; no filtra tokens ni abre conexiones."""
    profiles = (
        ("github", ("GITHUB_TOKEN",), "GitHub API; el token no prueba acceso al repositorio."),
        ("supabase", ("BAQUEANO_SUPABASE_URL", "BAQUEANO_SUPABASE_PUBLISHABLE_KEY"),
         "Las claves publicables no sustituyen permisos, RLS ni credenciales del proyecto."),
        ("vercel", ("VERCEL_TOKEN",), "Requiere cuenta, proyectos y permisos autorizados."),
        ("home_assistant", ("SKY_HA_URL", "SKY_HA_TOKEN"), "Dispositivos con autorización del propietario."),
        ("whatsapp_meta", ("SKY_WA_TOKEN", "SKY_WA_PHONE_ID"), "API oficial, número y permisos activos."),
    )
    return {
        "connectors": [
            {"name": name, "credentials_declared": all(bool(os.getenv(key, "").strip()) for key in keys),
             "connection_verified": False, "access_verified": False, "note": note}
            for name, keys, note in profiles
        ],
        "mcp_protocol_active": False,
        "note": "Inventario para preparar conectores; no finge servidores MCP ni ejecuta operaciones remotas.",
    }


def extract_document(filename: str, data: bytes) -> dict:
    """Texto digital con referencias trazables. Sin OCR ni guardado de documentos."""
    if not isinstance(filename, str) or not re.fullmatch(r"[\w .()\-]{1,100}\.(pdf|docx|txt|md|rst)", filename, re.I):
        raise ValueError("Formato permitido: PDF, DOCX, TXT, MD o RST.")
    if not isinstance(data, bytes) or not (0 < len(data) <= MAX_DOC_BYTES):
        raise ValueError("Documento vacío o mayor a 5 MB.")
    ext = filename.rsplit(".", 1)[-1].lower()
    source_id = hashlib.sha256(data).hexdigest()
    pages = []
    truncated = False
    if ext == "pdf":
        if not data.startswith(b"%PDF-"):
            raise ValueError("El archivo no contiene una cabecera PDF.")
        from pypdf import PdfReader
        try:
            reader = PdfReader(io.BytesIO(data), strict=False)
            if reader.is_encrypted:
                raise ValueError("PDF cifrado: se requiere autorización y contraseña.")
            truncated = len(reader.pages) > MAX_PAGES
            for number, page in enumerate(reader.pages[:MAX_PAGES], 1):
                txt = (page.extract_text() or "").strip()
                if txt:
                    pages.append({"page": number, "text": txt[:MAX_EXCERPT],
                                  "text_truncated": len(txt) > MAX_EXCERPT})
        except ValueError:
            raise
        except Exception as exc:
            raise ValueError("No se pudo leer el texto digital del PDF.") from exc
    elif ext == "docx":
        from docx import Document
        try:
            doc = Document(io.BytesIO(data))
            # Un DOCX no contiene paginación fija; no inventar números de página.
            text = "\n".join(p.text for p in doc.paragraphs if p.text.strip())
            for tab in doc.tables:
                for row in tab.rows:
                    text += "\n" + " | ".join(cell.text.strip() for cell in row.cells)
        except Exception as exc:
            raise ValueError("No se pudo leer el documento DOCX.") from exc
        truncated = len(text) > MAX_TOTAL_CHARS
        pages = [{"page": None, "text": text[:MAX_TOTAL_CHARS],
                  "text_truncated": truncated, "location": "documento_sin_paginacion_fija"}] if text.strip() else []
    else:
        text = data.decode("utf-8-sig", errors="replace")
        truncated = len(text) > MAX_TOTAL_CHARS
        pages = [{"page": None, "text": text[:MAX_TOTAL_CHARS],
                  "text_truncated": truncated, "location": "texto_sin_paginacion"}] if text.strip() else []
    if not pages:
        raise ValueError("No hay texto extraíble. Se requiere OCR explícito para escaneos.")
    return {
        "filename": filename,
        "source_sha256": source_id,
        "source_type": "user_document",
        "pages": pages,
        "page_count_extracted": len(pages),
        "truncated": truncated,
        "content_untrusted": True,
        "claims_verified": False,
        "no_persistent_copy": True,
        "note": "La página solo se cita si procede del PDF. Citas por ubicación no equivalen a verificación de hechos.",
    }


def _normalize(text: str) -> str:
    return re.sub(r"\s+", " ", str(text or "").casefold()).strip()


def crosscheck_claim(claim: str, sources: list[dict]) -> dict:
    """Localiza una cita literal presente; NUNCA declara verdadera una afirmación."""
    if not isinstance(claim, str) or not (8 <= len(claim.strip()) <= 1_200):
        raise ValueError("La afirmación debe tener entre 8 y 1200 caracteres.")
    if not isinstance(sources, list) or not 1 <= len(sources) <= 8:
        raise ValueError("Se requieren de 1 a 8 fuentes.")
    evidence = []
    for s in sources:
        if not isinstance(s, dict):
            raise ValueError("Fuente inválida.")
        source_id = s.get("source_sha256", "")
        page = s.get("page")
        quoted = s.get("quote", "")
        original = s.get("source_text", "")
        if not isinstance(source_id, str) or not re.fullmatch(r"[a-f0-9]{64}", source_id):
            raise ValueError("La fuente necesita un SHA256 de origen.")
        if not isinstance(quoted, str) or not 8 <= len(quoted.strip()) <= 450:
            raise ValueError("La cita tiene que contener entre 8 y 450 caracteres.")
        if not isinstance(original, str) or len(original) > MAX_TOTAL_CHARS:
            raise ValueError("Texto original de fuente inválido.")
        if page is not None and (type(page) is not int or page < 1 or page > 50_000):
            raise ValueError("Página no válida.")
        located = _normalize(quoted) in _normalize(original) if original else False
        evidence.append({"source_sha256": source_id, "page": page,
                         "quote": quoted if located else None,
                         "quote_located": located,
                         "citation": f"sha256:{source_id[:12]}" + (f":p{page}" if page else ":sin-pagina"),
                         "evidence_status": "text_located_not_claim_verified" if located else "quote_not_located"})
    found = [x for x in evidence if x["quote_located"]]
    return {
        "claim": claim,
        "evidence": evidence,
        "located_quotes": len(found),
        "distinct_source_ids_with_quotes": len({x["source_sha256"] for x in found}),
        "claim_verified": False,
        "cross_source_agreement_verified": False,
        "requires_human_review": True,
        "note": "Una coincidencia literal no demuestra verdad, causalidad ni acuerdo entre fuentes.",
    }


def formative_review(objective: str, criteria: list[str], student_response: str) -> dict:
    """Retroalimentación sin guardar datos, nombres ni calificar automáticamente."""
    if not isinstance(objective, str) or not 5 <= len(objective.strip()) <= 300:
        raise ValueError("Objetivo de aprendizaje inválido.")
    if not isinstance(criteria, list) or not 1 <= len(criteria) <= 6 or any(
            not isinstance(c, str) or not 5 <= len(c.strip()) <= 150 for c in criteria):
        raise ValueError("Definí entre 1 y 6 criterios formativos.")
    if not isinstance(student_response, str) or not 2 <= len(student_response.strip()) <= 800:
        raise ValueError("La respuesta del estudiante debe tener de 2 a 800 caracteres.")
    # Sin almacenamiento, puntuación algorítmica ni retorno del texto del estudiante.
    return {
        "objective": objective,
        "criteria": [{"criterion": c, "status": "review_required",
                      "question": f"¿Qué evidencia de la respuesta demuestra: {c}?"} for c in criteria],
        "feedback": "Gracias por tu respuesta. Revisemos juntos las evidencias de cada criterio y qué mejorar.",
        "grade": None,
        "correctness_graded": False,
        "student_answer_echoed": False,
        "student_response_persisted": False,
        "automatic_profile_created": False,
        "teacher_review_required": True,
        "privacy_note": "No introducir nombres, identificadores ni datos sensibles de estudiantes.",
    }


def kpi_snapshot(period: str = "hoy") -> dict:
    if period not in {"hoy", "semana"}:
        raise ValueError("Periodo desconocido.")
    from . import sky_metricas
    report = sky_metricas.resumen(period)
    model = report.get("modelo", {})
    return {
        "period": period,
        "sample_event_count": report.get("eventos", 0),
        "actions": report.get("acciones", {}),
        "voice": report.get("voz", {}),
        "api": report.get("api", {}),
        "model": model,
        "service": report.get("servicio", {}),
        "slo": report.get("slo", []),
        "cost_usd": (model.get("cost_usd") if isinstance(model, dict) else None),
        "cost_verified": bool(isinstance(model, dict) and model.get("cost_usd") is not None
                              and model.get("cost_estimated") is False),
        "null_means_no_observed_data": True,
        "note": "Datos del registro SKY; sin muestras se devuelven nulos, no números inventados.",
    }


def subagent_preview(task: str, selected_skills: list[str] | None = None) -> dict:
    """Plan de competencias, SIN ejecutar ni delegar al sistema del propietario."""
    from . import sky_pro_skills, sky_subagents
    proposal = sky_pro_skills.skill_plan(task, selected_skills)
    return {
        "central_agent": "SKY",
        "task": task,
        "skills": proposal["skills"],
        "available_roles": sky_subagents.catalogue()["roles"],
        "subagents_executed": 0,
        "actions_executed": 0,
        "owner_approval_required": True,
        "can_proceed_via": "existing_SKY_workflows_after_authorization",
        "note": "No se crea un subagente ni se ejecutan herramientas con una vista previa.",
    }
