"""SKY Office Studio: creación REAL de documentos, libros, diapositivas y PDF.

Generación local de archivos editables usando contenido proporcionado o texto de un
modelo elegido explícitamente por el usuario. Sin subida automática a terceros.
"""
from __future__ import annotations
from dataclasses import dataclass
from io import BytesIO, StringIO
from pathlib import Path
import csv
import re
import tempfile
import html
import unicodedata

FORMATS = {
    'docx': ('Word editable', 'application/vnd.openxmlformats-officedocument.wordprocessingml.document'),
    'xlsx': ('Excel editable', 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'),
    'pptx': ('PowerPoint editable', 'application/vnd.openxmlformats-officedocument.presentationml.presentation'),
    'pdf': ('PDF', 'application/pdf'),
    'csv': ('CSV', 'text/csv; charset=utf-8'),
    'md': ('Markdown', 'text/markdown; charset=utf-8'),
    'html': ('HTML', 'text/html; charset=utf-8'),
}
MAX_CONTENT = 50_000

@dataclass(frozen=True)
class GeneratedFile:
    filename: str
    data: bytes
    media_type: str
    description: str


def _clean(text: object, limit: int=MAX_CONTENT) -> str:
    if text is None: return ''
    if not isinstance(text,str):raise ValueError('El contenido debe ser texto.')
    if len(text)>limit:raise ValueError('Texto demasiado largo para un documento.')
    return text.replace('\x00','').strip()


def _safe_name(name: str) -> str:
    slug=re.sub(r'[^a-z0-9_-]+','_',unicodedata.normalize('NFKD',name).encode('ascii','ignore').decode().casefold().strip())[:70].strip('_')
    return slug or 'documento_sky'


def _segments(content: str, title: str):
    """Extrae secciones desde Markdown elemental sin interpretar HTML ni código."""
    lines=content.splitlines()
    sections=[]
    current_title='Contenido'
    parts=[]
    for raw in lines:
        line=raw.strip()
        match=re.match(r'^#{1,3}\s+(.+)$',line)
        if match:
            if parts:sections.append((current_title,'\n'.join(parts).strip()));parts=[]
            current_title=match.group(1)[:120]
        else:parts.append(raw)
    if parts:sections.append((current_title,'\n'.join(parts).strip()))
    if not sections:sections=[('Contenido',content or f'Documento: {title}')]
    return sections[:32]


def _sanitize_cells(headers,rows):
    if len(headers)>40 or len(rows)>300:raise ValueError('Máximo 40 columnas y 300 filas.')
    cols=[_clean(str(h),100) or f'Columna {i+1}' for i,h in enumerate(headers)]
    if not cols:raise ValueError('Se requiere al menos una columna.')
    matrix=[]
    for row in rows:
        if not isinstance(row,(tuple,list)) or len(row)>len(cols):raise ValueError('Fila de tabla inválida.')
        vals=[]
        for v in row:
            if v is None or isinstance(v,(bool,int,float)):vals.append(v)
            elif isinstance(v,str):
                if len(v)>1000:raise ValueError('El valor de una celda supera 1000 caracteres.')
                # Cifras CSV deben seguir siendo números para Excel, sin perder ceros iniciales.
                normalized=v.strip()
                if re.fullmatch(r'-?(?:0|[1-9]\d{0,12})',normalized):vals.append(int(normalized))
                elif re.fullmatch(r'-?(?:0|[1-9]\d{0,12})\.\d{1,8}',normalized):vals.append(float(normalized))
                else:
                    # Nunca evaluar texto externo como fórmula; el escritor exporta tipo texto.
                    vals.append("'"+v if v.lstrip().startswith(('=','+','-','@')) else v)
            else:raise ValueError('Tipo de celda no permitido.')
        matrix.append(vals+[None]*(len(cols)-len(vals)))
    return cols,matrix


def _table_from_csv(text):
    parsed=list(csv.reader(StringIO(text)))
    if not parsed:return [],[]
    return parsed[0],parsed[1:]


def _dataset(spec):
    headers=spec.get('headers') or []
    rows=spec.get('rows') or []
    if not isinstance(headers,list) or not isinstance(rows,list):
        raise ValueError('Encabezados y filas deben ser listas.')
    if not headers and spec.get('csv'):
        headers,rows=_table_from_csv(_clean(spec['csv'],30_000))
    if not headers:
        headers=['Elemento','Descripción','Estado']
        rows=[['Registro 1','Editable en Excel','Pendiente'],['Registro 2','','Pendiente']]
    return _sanitize_cells(headers,rows)


def _docx(title,sections,headers,rows):
    from docx import Document
    from docx.shared import Inches,Pt,RGBColor
    doc=Document()
    sec=doc.sections[0];sec.top_margin=Inches(.80);sec.bottom_margin=Inches(.75)
    sec.left_margin=Inches(.90);sec.right_margin=Inches(.90)
    normal=doc.styles['Normal'];normal.font.name='Aptos';normal.font.size=Pt(10.5)
    normal.paragraph_format.space_after=Pt(7)
    for label,size in [('Title',25),('Heading 1',16)]:
        s=doc.styles[label];s.font.name='Aptos Display';s.font.size=Pt(size);s.font.color.rgb=RGBColor(20,81,110)
    doc.add_heading(title,0)
    doc.add_paragraph('Preparado con SKY Office Studio',style='Subtitle')
    for heading,body in sections:
        doc.add_heading(heading,level=1)
        for line in body.splitlines():
            if not line.strip():continue
            if line.lstrip().startswith(('- ','* ','• ')):doc.add_paragraph(line.lstrip()[2:],style='List Bullet')
            else:doc.add_paragraph(line)
    if headers and rows:
        doc.add_heading('Datos',level=1)
        table=doc.add_table(rows=1,cols=len(headers));table.style='Light Shading Accent 1'
        for c,val in zip(table.rows[0].cells,headers):c.text=str(val)
        for row in rows:
            cells=table.add_row().cells
            for cell,val in zip(cells,row):cell.text='' if val is None else str(val)
    buf=BytesIO();doc.save(buf);return buf.getvalue()


def _pptx(title,sections,headers,rows):
    from pptx import Presentation
    from pptx.util import Inches,Pt
    from pptx.dml.color import RGBColor
    from pptx.enum.text import PP_ALIGN
    p=Presentation();p.slide_width=Inches(13.333);p.slide_height=Inches(7.5)
    NAVY=RGBColor(10,21,42); CYAN=RGBColor(40,207,221); WHITE=RGBColor(238,246,255)
    def base(slide):
        slide.background.fill.solid();slide.background.fill.fore_color.rgb=NAVY
        rule=slide.shapes.add_shape(1, Inches(.65), Inches(6.95), Inches(12), Inches(.025))
        rule.fill.solid();rule.fill.fore_color.rgb=CYAN;rule.line.fill.background()
    def txt(slide,text,x,y,w,h,size=24,color=WHITE,bold=False):
        box=slide.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h))
        tf=box.text_frame;tf.word_wrap=True
        for i,line in enumerate(text.splitlines() or ['']):
            par=tf.paragraphs[0] if i==0 else tf.add_paragraph()
            par.text=line[:210];par.font.size=Pt(size);par.font.bold=bold;par.font.color.rgb=color
        return box
    slide=p.slides.add_slide(p.slide_layouts[6]);base(slide)
    txt(slide,'SKY OFFICE STUDIO',.85,.6,8,.45,15,CYAN,True)
    txt(slide,title, .85,2.1,11.6,2.3,39,WHITE,True)
    txt(slide,'Presentación editable · creada con SKY',.90,5.7,10,.65,17,CYAN)
    for idx,(heading,body) in enumerate(sections[:24],1):
        slide=p.slides.add_slide(p.slide_layouts[6]);base(slide)
        txt(slide,heading,.8,.55,11.5,1.1,29,CYAN,True)
        bullets=[ln.lstrip(' -*•\t') for ln in body.splitlines() if ln.strip()]
        bullets=bullets[:7] or ['Contenido editable']
        for j,bullet in enumerate(bullets):
            txt(slide,'• '+bullet[:160],1.05,1.7+j*.68,11.2,.6,20)
        txt(slide,f'SKY IA  •  {idx:02}',.8,7.08,6,.25,10,CYAN)
    if headers and rows:
        slide=p.slides.add_slide(p.slide_layouts[6]);base(slide)
        txt(slide,'Datos y cifras',.8,.55,11,1.0,29,CYAN,True)
        cols=min(len(headers),6);rs=min(len(rows),7)
        table=slide.shapes.add_table(rs+1,cols,Inches(.7),Inches(1.75), Inches(11.9), Inches(3.95)).table
        for j,h in enumerate(headers[:cols]):table.cell(0,j).text=str(h)[:36]
        for i,row in enumerate(rows[:rs],1):
            for j,v in enumerate(row[:cols]):table.cell(i,j).text=str(v or '')[:50]
        for row in table.rows:
            for cell in row.cells:
                cell.fill.solid();cell.fill.fore_color.rgb=RGBColor(26,48,73)
                for par in cell.text_frame.paragraphs:
                    par.font.color.rgb=WHITE;par.font.size=Pt(12)
    b=BytesIO();p.save(b);return b.getvalue()


def _pdf(title,sections,headers,rows):
    from reportlab.platypus import SimpleDocTemplate,Paragraph,Spacer,Table,TableStyle,KeepTogether
    from reportlab.lib import colors
    from reportlab.lib.styles import getSampleStyleSheet,ParagraphStyle
    from reportlab.lib.enums import TA_LEFT
    from reportlab.lib.pagesizes import A4
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.lib.units import cm
    font='Helvetica'
    for path in ('/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf','/usr/share/fonts/truetype/liberation2/LiberationSans-Regular.ttf'):
        if Path(path).is_file():
            try:pdfmetrics.registerFont(TTFont('SKYUnicode',path));font='SKYUnicode';break
            except Exception:pass
    styles=getSampleStyleSheet()
    styles.add(ParagraphStyle(name='SKYTitle',parent=styles['Title'],fontName=font,fontSize=22,leading=28,textColor=colors.HexColor('#154d6e'),spaceAfter=22))
    styles.add(ParagraphStyle(name='SKYHeading',parent=styles['Heading2'],fontName=font,fontSize=13,leading=18,textColor=colors.HexColor('#154d6e'),spaceBefore=13,spaceAfter=8))
    styles.add(ParagraphStyle(name='SKYBody',parent=styles['BodyText'],fontName=font,fontSize=10,leading=15,spaceAfter=7))
    def esc(t):return html.escape(str(t)).replace('\n','<br/>')
    story=[Paragraph(esc(title),styles['SKYTitle']),Paragraph('SKY Office Studio',styles['SKYBody']),Spacer(1,15)]
    for heading,body in sections:
        story.append(Paragraph(esc(heading),styles['SKYHeading']))
        for line in body.splitlines():
            if line.strip():story.append(Paragraph(esc(line),styles['SKYBody']))
    if headers and rows:
        story.append(Paragraph('Datos',styles['SKYHeading']))
        data=[[Paragraph(esc(v),styles['SKYBody']) for v in r] for r in [headers]+rows[:50]]
        widths=[(A4[0]-3*cm)/len(headers)]*len(headers)
        t=Table(data,colWidths=widths,repeatRows=1,hAlign='LEFT')
        t.setStyle(TableStyle([('BACKGROUND',(0,0),(-1,0),colors.HexColor('#154d6e')),('TEXTCOLOR',(0,0),(-1,0),colors.white),('GRID',(0,0),(-1,-1),.35,colors.HexColor('#CBD5E1')),('VALIGN',(0,0),(-1,-1),'TOP'),('ROWBACKGROUNDS',(0,1),(-1,-1),[colors.whitesmoke,colors.white])]))
        story.append(t)
    buf=BytesIO()
    SimpleDocTemplate(buf,pagesize=A4,rightMargin=1.5*cm,leftMargin=1.5*cm,topMargin=1.6*cm,bottomMargin=1.7*cm,
                      title=title,author='SKY Office Studio').build(story)
    return buf.getvalue()


def _xlsx(title,headers,rows,template):
    try:from artifact_tool import Workbook, SpreadsheetFile
    except ImportError:return _xlsx_portable(title,headers,rows,template)
    wb=Workbook.create()
    sheet=wb.worksheets.add('Datos')
    sheet.get_range('A1:D1').merge();sheet.get_range('A1').values=[[title]]
    sheet.get_range('A1').format.font.bold=True
    sheet.get_range('A1').format.font.color='#FFFFFF'
    sheet.get_range('A1:D1').format.fill='#12354F'
    sheet.get_range('A1:D1').format.row_height=30
    if template=='presupuesto' and not rows:
        headers=['Concepto','Presupuesto','Gasto real','Diferencia']
        rows=[[f'Actividad {i}',0,0,None] for i in range(1,7)]
        matrix=[headers]+rows+[['TOTAL',None,None,None]]
        sheet.get_range('A3:D10').values=matrix
        for i in range(4,10):sheet.get_range(f'D{i}').formulas=[[f'=B{i}-C{i}']]
        for col in 'BCD':sheet.get_range(f'{col}10').formulas=[[f'=SUM({col}4:{col}9)']]
        sheet.get_range('B4:D10').set_number_format('#,##0.00')
    else:
        if len(headers)>26:raise ValueError('Excel acepta hasta 26 columnas en esta versión.')
        headers=list(headers)
        sheet.get_range_by_indexes(2,0,1,len(headers)).values=[headers]
        if rows:sheet.get_range_by_indexes(3,0,len(rows),len(headers)).values=rows
    sheet.get_range_by_indexes(2,0,1,len(headers)).format={
        'fill':'#155E75','font':{'bold':True,'color':'#FFFFFF'},'row_height':26}
    for col in range(len(headers)):
        sheet.get_range_by_indexes(2,col, max(len(rows)+2,2),1).format.column_width = (24 if col==0 else 19)
    sheet.freeze_panes.freeze_rows(3)
    with tempfile.TemporaryDirectory() as tmp:
        path=Path(tmp)/'sky.xlsx'
        SpreadsheetFile.export_xlsx(wb).save(str(path))
        return path.read_bytes()


def _xlsx_portable(title,headers,rows,template):
    """Motor Excel portátil, basado en dependencia pública (sin librería privada)."""
    try:import xlsxwriter
    except ImportError as exc:raise RuntimeError('Instalá xlsxwriter para crear archivos Excel.') from exc
    buf=BytesIO()
    wb=xlsxwriter.Workbook(buf,{'in_memory':True})
    ws=wb.add_worksheet('Datos')
    title_style=wb.add_format({'bg_color':'#12354F','font_color':'#FFFFFF','bold':True,'font_size':17,'valign':'vcenter'})
    header_style=wb.add_format({'bg_color':'#155E75','font_color':'#FFFFFF','bold':True,'text_wrap':True})
    money_style=wb.add_format({'num_format':'#,##0.00'})
    ws.set_row(0,35)
    ws.merge_range(0,0,0,3,title,title_style)
    if template=='presupuesto' and not rows:
        headers=['Concepto','Presupuesto','Gasto real','Diferencia']
        rows=[[f'Actividad {i}',0,0,None] for i in range(1,7)]
        for j,h in enumerate(headers):ws.write(2,j,h,header_style)
        for i,row in enumerate(rows,3):
            ws.write_string(i,0,row[0]);ws.write_number(i,1,0,money_style);ws.write_number(i,2,0,money_style)
            ws.write_formula(i,3,f'=B{i+1}-C{i+1}',money_style)
        ws.write_string(9,0,'TOTAL',header_style)
        for j,col in enumerate('BCD',1):ws.write_formula(9,j,f'=SUM({col}4:{col}9)',money_style)
    else:
        for j,h in enumerate(headers):ws.write(2,j,h,header_style)
        for i,row in enumerate(rows,3):
            for j,v in enumerate(row):
                if isinstance(v,str):
                    # write_string evita la ejecución de expresiones externas.
                    ws.write_string(i,j,v.lstrip("'") if v.startswith("'") else v)
                elif isinstance(v,(int,float)) and not isinstance(v,bool):ws.write_number(i,j,v)
                elif v is not None:ws.write_boolean(i,j,bool(v))
    for col in range(len(headers)):ws.set_column(col,col,27 if col==0 else 20)
    ws.freeze_panes(3,0)
    wb.close()
    return buf.getvalue()


def create_document(spec:dict)->GeneratedFile:
    """Genera archivos binarios verídicos a partir de una especificación limitada."""
    if not isinstance(spec,dict):raise ValueError('Se esperaba una especificación JSON.')
    kind=str(spec.get('kind','docx')).lower()
    if kind not in FORMATS:raise ValueError('Formato no permitido.')
    title=_clean(spec.get('title') or 'Mi creación SKY',160)
    if not title:title='Mi creación SKY'
    content=_clean(spec.get('content') or spec.get('prompt') or '')
    if not content:content='Escribí aquí tu contenido y editalo según lo necesités.'
    headers,rows=([],[])
    if kind=='xlsx' and str(spec.get('template','general')).lower()=='presupuesto' and not spec.get('headers') and not spec.get('csv'):
        headers=['Concepto','Presupuesto','Gasto real','Diferencia'];rows=[]
    elif kind in ('xlsx','csv') or spec.get('headers') or spec.get('csv'):
        headers,rows=_dataset(spec)
    sections=_segments(content,title)
    template=str(spec.get('template','general')).lower()
    if template not in ('general','presupuesto'):raise ValueError('Plantilla no disponible.')
    if kind=='docx':data=_docx(title,sections,headers,rows)
    elif kind=='pptx':data=_pptx(title,sections,headers,rows)
    elif kind=='pdf':data=_pdf(title,sections,headers,rows)
    elif kind=='xlsx':data=_xlsx(title,headers,rows,template)
    elif kind=='csv':
        buf=StringIO();writer=csv.writer(buf);writer.writerow(headers);writer.writerows(rows)
        data=('\ufeff'+buf.getvalue()).encode('utf-8')
    elif kind=='md':data=(f'# {title}\n\n'+content+'\n').encode('utf-8')
    else:data=(f'<!doctype html><html lang="es"><head><meta charset="utf-8"><title>{html.escape(title)}</title></head><body><h1>{html.escape(title)}</h1><main><pre style="white-space:pre-wrap">{html.escape(content)}</pre></main></body></html>').encode('utf-8')
    label,mime=FORMATS[kind]
    return GeneratedFile(f'{_safe_name(title)}.{kind}',data,mime,label)
